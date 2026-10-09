from __future__ import annotations

import hashlib
import json
from datetime import datetime

import pytest
from openpyxl import Workbook, load_workbook

from budget_dashboard.cli import main
from budget_dashboard.workbook import (
    _read_budget_grid,
    _read_journal,
    build_dashboard_data,
    build_dashboard_payload,
)
from finance_hub.budget_api import BudgetPayloadCache


def _workbook(path):
    wb = Workbook()
    journal = wb.active
    journal.title = "Journal"
    journal.append(["Journal"])
    journal.append(["Date", "Debit", "Credit", "Amount", "Note"])
    journal.append([datetime(2026, 1, 5), "Primary Checking", "Salary", 1000, "Pay"])
    journal.append([datetime(2026, 1, 8), "Food", "Primary Checking", 200, "Groceries"])
    journal.append([datetime(2026, 1, 9), "Investment Account", "Primary Checking", 100, "Savings"])
    journal.append([datetime(2026, 1, 31), "Investment Account", "Investment Income", 50, "Month close"])
    journal.append([datetime(2026, 2, 5), "Primary Checking", "Salary", 1200, "Pay"])
    journal.append([datetime(2026, 2, 8), "Electricity", "Primary Checking", 50, "Power"])

    budget = wb.create_sheet("Budget")
    budget.cell(13, 2, datetime(2026, 1, 31))
    budget.cell(13, 3, datetime(2026, 2, 28))
    labels = {16: "Salary", 27: "Utilities", 28: "Food"}
    for row, label in labels.items():
        budget.cell(row, 1, label)
    budget.cell(16, 2, 900)
    budget.cell(16, 3, 1100)
    budget.cell(27, 2, -60)
    budget.cell(27, 3, -60)
    budget.cell(28, 2, -250)
    # Food's February cell is deliberately left blank (no assertions elsewhere
    # depend on it) so _read_budget_grid's None-for-blank-cell behavior has a
    # real blank cell inside an existing row to exercise.

    ledger = wb.create_sheet("Ledger")
    ledger.append(["Ledger"])
    ledger.append(["Account", "Beginning Balance"])
    ledger.append(["Primary Checking", 100])
    ledger.append(["Investment Account", 0])
    wb.save(path)


def test_payload_matches_journal_and_budget_without_changing_workbook(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)
    before = hashlib.sha256(path.read_bytes()).hexdigest()

    data = build_dashboard_payload(path)

    assert hashlib.sha256(path.read_bytes()).hexdigest() == before
    assert data["latest_period"] == "2026-02"
    assert "transactions" not in data
    assert data["monthly"]["2026-01"]["income"] == 1000
    assert data["monthly"]["2026-01"]["spending"] == 200
    assert data["monthly"]["2026-01"]["spending_budget"] == 310
    assert data["monthly"]["2026-01"]["savings_rate"] == 0.1
    assert data["monthly"]["2026-01"]["investment_income"] == 50
    investment = next(
        row for row in data["monthly"]["2026-01"]["categories"]
        if row["kind"] == "investment"
    )
    assert investment == {
        "category": "Investment Income", "kind": "investment",
        "actual": 50, "budget": None, "variance": None,
    }
    assert data["ytd"]["2026-02"]["income"] == 2200
    utility = next(
        row for row in data["monthly"]["2026-02"]["categories"]
        if row["category"] == "Utilities"
    )
    assert utility["actual"] == 50
    assert data["balances"]["2026-01"]["net_worth"] == 950


def test_dashboard_data_retains_read_only_transaction_rows_separately(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)

    payload, transactions = build_dashboard_data(path)

    assert "transactions" not in payload
    assert transactions[0] == {
        "row_number": 3,
        "date": "2026-01-05",
        "period": "2026-01",
        "debit": "Primary Checking",
        "credit": "Salary",
        "amount": 1000.0,
        "note": "Pay",
    }


def test_transaction_query_filters_pages_and_calculates_dashboard_impact(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)
    cache = BudgetPayloadCache(path)

    food = cache.query_transactions(
        start="2026-01-01", end="2026-01-31", category="Food", limit=1
    )
    assert food["total"] == 1
    assert food["items"][0]["row_number"] == 4
    assert food["items"][0]["impact_cad"] == 200

    cash = cache.query_transactions(account="Primary Checking", q="groceries")
    assert cash["total"] == 1
    assert cash["items"][0]["impact_cad"] == -200

    income = cache.query_transactions(kind="income", start="2026-01-01", end="2026-01-31")
    assert income["total"] == 1
    assert income["items"][0]["impact_cad"] == 1000

    second_page = cache.query_transactions(limit=2, offset=2)
    assert second_page["limit"] == 2
    assert second_page["offset"] == 2
    assert len(second_page["items"]) == 2


@pytest.mark.parametrize(
    "kwargs",
    [
        {"start": "not-a-date"},
        {"start": "2026-02-01", "end": "2026-01-01"},
        {"limit": 0},
        {"limit": 501},
        {"offset": -1},
    ],
)
def test_transaction_query_rejects_invalid_filters(tmp_path, kwargs):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)

    with pytest.raises(ValueError):
        BudgetPayloadCache(path).query_transactions(**kwargs)


def test_cli_writes_dashboard_payload_as_json(tmp_path):
    workbook = tmp_path / "Personal Budget.xlsx"
    output = tmp_path / "dashboard.json"
    _workbook(workbook)

    assert main(["--workbook", str(workbook), "--output", str(output)]) == 0
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["latest_period"] == "2026-02"


def test_investment_losses_remain_separate_from_operating_income(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)
    workbook = load_workbook(path)
    workbook["Journal"].append([
        datetime(2026, 2, 28), "Investment Income", "Investment Account", 25,
        "Month close",
    ])
    workbook.save(path)
    workbook.close()

    february = build_dashboard_payload(path)["monthly"]["2026-02"]

    assert february["income"] == 1200
    assert february["investment_income"] == -25
    investment = next(row for row in february["categories"] if row["kind"] == "investment")
    assert investment["actual"] == -25
    assert investment["budget"] is None
    assert investment["variance"] is None


def test_read_journal_returns_signed_actuals_and_transactions(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)

    journal_data = _read_journal(wb)
    wb.close()

    assert journal_data.transactions[0]["debit"] == "Primary Checking"
    assert journal_data.actuals[("2026-01", "Salary")] == 1000.0
    assert journal_data.actuals[("2026-01", "Food")] == -200.0
    assert journal_data.account_debits[("2026-01", "Primary Checking")] == 1000.0
    assert journal_data.savings_by_month["2026-01"] == 100.0


def test_read_journal_recognizes_employer_neutral_espp_account(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)
    workbook = load_workbook(path)
    workbook["Journal"].append([
        datetime(2026, 2, 10), "Employer ESPP", "Cash", 75, "Employer ESPP",
    ])
    workbook.save(path)
    workbook.close()

    workbook = load_workbook(path, read_only=True, data_only=True)
    journal_data = _read_journal(workbook)
    workbook.close()

    assert journal_data.espp_by_month["2026-02"] == 75.0


def test_read_budget_grid_preserves_none_for_blank_cells(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _workbook(path)
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)

    budget_by_month = _read_budget_grid(wb)
    wb.close()

    assert budget_by_month["2026-01"]["Salary"] == 900.0
    assert budget_by_month["2026-01"]["Utilities"] == -60.0
    # "Rent" has no row in the fixture's Budget sheet at all -> not a key.
    assert "Rent" not in budget_by_month["2026-01"]
    # "Food" has a row, but its February cell is left blank in the fixture ->
    # None, not 0.0, distinguishing "no entry yet" from "entered as zero".
    assert budget_by_month["2026-02"]["Food"] is None
