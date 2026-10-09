"""Account/category names for the journal-entry UI, read from Ledger!A:A.

Reuses the category/account constants budget_dashboard already maintains
rather than defining a second copy of the chart of accounts.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from budget_dashboard.workbook import (
    ASSET_ACCOUNTS,
    EXPENSE_CATEGORIES,
    INCOME_CATEGORIES,
    LIABILITY_ACCOUNTS,
)

_ASSET_SET = set(ASSET_ACCOUNTS)
_LIABILITY_SET = set(LIABILITY_ACCOUNTS)
_INCOME_SET = set(INCOME_CATEGORIES)
_EXPENSE_SET = set(EXPENSE_CATEGORIES)


def read_ledger_accounts(workbook_path: str | Path) -> list[str]:
    """Every name in Ledger!A:A, in sheet order. Drives Debit/Credit dropdowns."""
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        ledger = wb["Ledger"]
        return [
            str(row[0]).strip()
            for row in ledger.iter_rows(min_row=3, max_col=1, values_only=True)
            if row[0] not in (None, "")
        ]
    finally:
        wb.close()


def account_kind(name: str) -> str:
    """Classify a Ledger name as 'asset', 'liability', 'income', 'expense', or 'other'."""
    if name in _ASSET_SET:
        return "asset"
    if name in _LIABILITY_SET:
        return "liability"
    if name in _INCOME_SET:
        return "income"
    if name in _EXPENSE_SET:
        return "expense"
    return "other"


def is_balance_account(name: str) -> bool:
    """True for accounts a bank CSV could plausibly represent (asset or liability)."""
    return name in _ASSET_SET or name in _LIABILITY_SET
