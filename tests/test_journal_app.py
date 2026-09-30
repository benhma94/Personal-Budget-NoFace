from __future__ import annotations

import json
from datetime import date

import pytest
from openpyxl import load_workbook

from journal_entry.app import JournalApp
from journal_entry.csv_import import CsvFormatError

from .journal_fixture import build_fixture_workbook

_ACCOUNTS = [
    "Cash", "Rewards Card", "Brokerage Cash", "Secondary Checking", "Credit Card",
    "Food", "Transit", "Salary", "Recurring Fees",
]

_CSV = """Date,Description,Amount
2026-08-01,EXAMPLE MARKET #1042,-41.20
2026-08-02,EXAMPLE MARKET #1042,-18.75
2026-08-03,EXAMPLE SERVICE,-28.25
"""


@pytest.fixture
def app(tmp_path):
    workbook_path = tmp_path / "Personal Budget.xlsx"
    build_fixture_workbook(workbook_path, _ACCOUNTS)
    (tmp_path / "merchant_seed_rules.json").write_text(
        json.dumps([["EXAMPLE SERVICE", "Recurring Fees", "Software"]]), encoding="utf-8"
    )
    instance = JournalApp.create(workbook_path, tmp_path / "journal_entry.sqlite")
    yield instance
    instance.close()


def test_accounts_reflects_ledger(app):
    result = app.accounts()
    assert "Rewards Card" in result["accounts"]
    assert result["kinds"]["Food"] == "expense"
    assert result["kinds"]["Rewards Card"] == "liability"


def test_sniff_reports_headers_and_no_known_mapping_first_time(app):
    result = app.sniff(_CSV)
    assert result["headers"] == ["Date", "Description", "Amount"]
    assert result["known_source"] is None


def test_sniff_guesses_date_column_and_format_for_new_layout(app):
    csv_text = "Transaction Date,Description,Amount\n08/25/2026,EXAMPLE MARKET,-41.20\n08/26/2026,CORNER STORE,-18.75\n"
    result = app.sniff(csv_text)
    assert result["guessed_date_col"] == "Transaction Date"
    assert result["guessed_date_format"] == "%m/%d/%Y"


def test_preview_dates_reports_matches_and_mismatches(app):
    csv_text = "Date,Description,Amount\n08/25/2026,EXAMPLE MARKET,-41.20\nnot-a-date,CORNER STORE,-18.75\n"
    result = app.preview_dates(csv_text, "Date", "%m/%d/%Y")
    assert result["samples"] == [
        {"raw": "08/25/2026", "parsed": "2026-08-25"},
        {"raw": "not-a-date", "parsed": None},
    ]


def test_preview_dates_empty_when_column_unset(app):
    result = app.preview_dates(_CSV, "Nope", "%Y-%m-%d")
    assert result["samples"] == []


def test_sniff_omits_guess_when_source_is_known(app):
    sniffed = app.sniff(_CSV)
    app.save_mapping(
        sniffed["header_signature"], "Rewards Card export", "Rewards Card",
        {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}, "%Y-%m-%d",
    )
    result = app.sniff(_CSV)
    assert result["known_source"] is not None
    assert result["guessed_date_col"] is None
    assert result["guessed_date_format"] is None


def test_import_without_mapping_raises(app):
    with pytest.raises(ValueError):
        app.import_csv(_CSV)


def test_save_mapping_then_import_reuses_it(app):
    sniffed = app.sniff(_CSV)
    app.save_mapping(
        sniffed["header_signature"], "Rewards Card export", "Rewards Card",
        {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}, "%Y-%m-%d",
    )
    result = app.import_csv(_CSV, header_signature=sniffed["header_signature"])
    assert result["read"] == 3
    assert result["new"] == 3
    assert result["duplicate"] == 0


def test_reimporting_same_file_is_all_duplicates(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    result = app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    assert result["new"] == 0
    assert result["duplicate"] == 3


def test_seed_rule_pre_fills_suggestion_without_confirming(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    service_row = next(t for t in app.list_transactions() if "EXAMPLE SERVICE" in t["description"])
    assert service_row["category"] == "Recurring Fees"
    assert service_row["status"] == "new"  # suggested, not confirmed


def test_categorize_confirms_and_learns_a_new_rule(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    example_market = [t for t in app.list_transactions() if "EXAMPLE MARKET" in t["description"]]
    fingerprints = [t["fingerprint"] for t in example_market]

    app.categorize(fingerprints, "Food", "Groceries")

    categorized = app.list_transactions(status="categorized")
    assert len(categorized) == 2
    assert all(t["category"] == "Food" for t in categorized)

    rule_patterns = {r["pattern"] for r in app.rules()}
    assert "EXAMPLE MARKET" in rule_patterns


def test_post_writes_journal_rows_and_marks_posted(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    fingerprints = [t["fingerprint"] for t in app.list_transactions()]
    app.categorize(fingerprints, "Food", "August groceries and subscriptions")

    preview = app.post_preview(date(2026, 8, 30))
    assert len(preview["lines"]) == 1
    assert preview["lines"][0]["amount"] == pytest.approx(88.20)

    result = app.post(date(2026, 8, 30))
    assert result["line_count"] == 1

    wb = load_workbook(app.workbook_path, read_only=True, data_only=True)
    row = list(wb["Journal"].iter_rows(min_row=result["first_row"], max_row=result["first_row"],
                                        max_col=5, values_only=True))[0]
    wb.close()
    assert row[1:3] == ("Food", "Rewards Card")
    assert row[3] == pytest.approx(88.20)

    posted = app.list_transactions(status="posted")
    assert len(posted) == 3


def test_post_with_nothing_categorized_raises(app):
    with pytest.raises(ValueError):
        app.post(date(2026, 8, 30))


def test_transfer_workflow_confirm_post_unlink(app):
    out_csv = "Date,Description,Amount\n2026-08-12,PAYMENT,-631.11\n"
    in_csv = "Date,Description,Amount\n2026-08-12,PAYMENT RECEIVED,631.11\n"
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(out_csv, account="Brokerage Cash", mapping=mapping, date_format="%Y-%m-%d")
    app.import_csv(in_csv, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")

    suggestions = app.suggested_transfers()
    assert len(suggestions) == 1
    outgoing_fp = suggestions[0]["outgoing"]["fingerprint"]
    incoming_fp = suggestions[0]["incoming"]["fingerprint"]

    app.confirm_transfer(outgoing_fp, incoming_fp, "Pay Bill")
    # Confirmed transfers are excluded from the normal review queue.
    assert app.list_transactions(status=["new", "categorized"]) == []

    preview = app.post_preview(date(2026, 8, 30))
    assert len(preview["lines"]) == 1
    assert preview["lines"][0]["debit"] == "Rewards Card"
    assert preview["lines"][0]["credit"] == "Brokerage Cash"

    app.unlink_transfer(outgoing_fp)
    assert len(app.list_transactions(status=["new", "categorized"])) == 2
    assert app.post_preview(date(2026, 8, 30))["lines"] == []


def test_import_bad_mapping_raises_csv_format_error(app):
    with pytest.raises(CsvFormatError):
        app.import_csv(
            _CSV, account="Rewards Card",
            mapping={"date_col": "Nope", "desc_col": "Description", "amount_col": "Amount"},
            date_format="%Y-%m-%d",
        )


def test_add_and_delete_rule(app):
    app.add_rule("COSTCO", "Shopping", "Costco run")
    assert any(r["pattern"] == "COSTCO" for r in app.rules())
    rule_id = next(r["id"] for r in app.rules() if r["pattern"] == "COSTCO")
    app.delete_rule(rule_id)
    assert not any(r["pattern"] == "COSTCO" for r in app.rules())


def test_undo_import_deletes_new_rows_and_skips_reviewed(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    result = app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    fingerprints = [row["fingerprint"] for row in result["rows"] if not row["duplicate"]]

    example_market_aug1 = next(
        t["fingerprint"] for t in app.list_transactions()
        if t["description"] == "EXAMPLE MARKET #1042" and t["txn_date"] == "2026-08-01"
    )
    app.categorize([example_market_aug1], "Food", "Groceries")

    undo_result = app.undo_import(fingerprints)

    assert undo_result == {"deleted": 2, "kept": 1}
    remaining = {t["fingerprint"] for t in app.list_transactions(status=["new", "categorized"])}
    assert remaining == {example_market_aug1}


def test_undo_import_dedupes_repeated_fingerprints(app):
    # Two identical rows in one CSV share a fingerprint; counting them twice
    # would report a phantom "already reviewed" row that was in fact deleted.
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    fp = app.list_transactions()[0]["fingerprint"]

    result = app.undo_import([fp, fp, fp])

    assert result == {"deleted": 1, "kept": 0}
    assert app.store.get_transaction(fp) is None


def test_batch_lines_reads_current_workbook_contents(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    fingerprints = [t["fingerprint"] for t in app.list_transactions()]
    app.categorize(fingerprints, "Food", "August groceries and subscriptions")
    result = app.post(date(2026, 8, 30))

    lines = app.batch_lines(result["batch_id"])

    assert len(lines) == 1
    assert lines[0]["row_number"] == result["first_row"]
    assert lines[0]["debit"] == "Food"
    assert lines[0]["credit"] == "Rewards Card"
    assert lines[0]["amount"] == pytest.approx(88.20)


def test_batch_lines_unknown_batch_raises(app):
    with pytest.raises(ValueError):
        app.batch_lines("nope")


def test_update_journal_row_corrects_a_posted_row(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    fingerprints = [t["fingerprint"] for t in app.list_transactions()]
    app.categorize(fingerprints, "Food", "August groceries and subscriptions")
    result = app.post(date(2026, 8, 30))
    original = app.batch_lines(result["batch_id"])[0]

    outcome = app.update_journal_row(
        original["row_number"], posting_date=date(2026, 8, 31), debit="Food",
        credit="Rewards Card", amount=100.0, note="Corrected", expected=original,
    )

    assert outcome["row_number"] == original["row_number"]
    updated = app.batch_lines(result["batch_id"])[0]
    assert updated["posting_date"] == "2026-08-31"
    assert updated["amount"] == 100.0
    assert updated["note"] == "Corrected"


def test_update_journal_row_outside_any_batch_raises(app):
    with pytest.raises(ValueError):
        app.update_journal_row(
            3, posting_date=date(2026, 8, 31), debit="Food", credit="Cash",
            amount=10.0, note="x", expected={
                "posting_date": "2026-08-30", "debit": "Food", "credit": "Cash",
                "amount": 10.0, "note": "x",
            },
        )


def test_update_journal_row_zero_amount_raises(app):
    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    fingerprints = [t["fingerprint"] for t in app.list_transactions()]
    app.categorize(fingerprints, "Food", "note")
    result = app.post(date(2026, 8, 30))
    original = app.batch_lines(result["batch_id"])[0]

    with pytest.raises(ValueError):
        app.update_journal_row(
            original["row_number"], posting_date=date(2026, 8, 31), debit="Food",
            credit="Rewards Card", amount=0.0, note="x", expected=original,
        )


def test_update_journal_row_stale_expected_raises_conflict_without_writing(app):
    from journal_entry.app import RowConflictError

    mapping = {"date_col": "Date", "desc_col": "Description", "amount_col": "Amount"}
    app.import_csv(_CSV, account="Rewards Card", mapping=mapping, date_format="%Y-%m-%d")
    fingerprints = [t["fingerprint"] for t in app.list_transactions()]
    app.categorize(fingerprints, "Food", "August groceries and subscriptions")
    result = app.post(date(2026, 8, 30))
    original = app.batch_lines(result["batch_id"])[0]
    stale_expected = {**original, "amount": original["amount"] + 1}

    with pytest.raises(RowConflictError):
        app.update_journal_row(
            original["row_number"], posting_date=date(2026, 8, 31), debit="Food",
            credit="Rewards Card", amount=100.0, note="Corrected", expected=stale_expected,
        )

    unchanged = app.batch_lines(result["batch_id"])[0]
    assert unchanged["amount"] == original["amount"]
