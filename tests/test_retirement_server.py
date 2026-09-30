from __future__ import annotations

import json
import threading
import urllib.error
import urllib.request
from datetime import datetime
from http.server import ThreadingHTTPServer

import pytest
from openpyxl import Workbook

from finance_hub.budget_api import BudgetPayloadCache
from finance_hub.idle_monitor import IdleMonitor
from finance_hub.portfolio_api import PortfolioRefreshJob
from finance_hub.server import make_handler


def _workbook(path):
    workbook = Workbook()
    journal = workbook.active
    journal.title = "Journal"
    journal.append(["Journal"])
    journal.append(["Date", "Debit", "Credit", "Amount", "Note"])
    journal.append([datetime(2026, 2, 28), "Food", "Cash", 590, "Groceries"])

    budget = workbook.create_sheet("Budget")
    budget.cell(13, 2, datetime(2026, 2, 28))
    budget.cell(28, 1, "Food")
    budget.cell(28, 2, -600)

    ledger = workbook.create_sheet("Ledger")
    ledger.append(["Ledger"])
    ledger.append(["Account", "Beginning Balance"])
    ledger.append(["Cash", 1000])
    workbook.save(path)


@pytest.fixture
def retirement_server(tmp_path):
    workbook_path = tmp_path / "Personal Budget.xlsx"
    _workbook(workbook_path)
    payload_path = tmp_path / "portfolio_payload.json"
    payload_path.write_text(json.dumps({
        "generated_at": "2026-02-28T12:00:00-05:00",
        "payload": {
            "as_of_date": "2026-02-28",
            "total_value_cad": 500_000,
            "performance": [{
                "Horizon": "Since-inception",
                "Return Basis": "Annualized",
                "Portfolio Return": 0.1383,
            }],
            "account_breakdown": [
                {"account_type": "TFSA", "value_cad": 180_000},
                {"account_type": "RRSP", "value_cad": 200_000},
            ],
            "account_positions": [
                {"account_type": "Non-registered", "total_acb": 70_000},
                {"account_type": "Non-registered", "total_acb": 5_000},
            ],
        },
    }), encoding="utf-8")
    budget_cache = BudgetPayloadCache(workbook_path)
    portfolio_job = PortfolioRefreshJob(tmp_path / "portfolio.xlsx", payload_path)
    handler = make_handler(
        object(),
        budget_cache,
        portfolio_job,
        payload_path,
        IdleMonitor(timeout_seconds=60),
        workbook_path,
    )
    httpd = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{httpd.server_address[1]}", payload_path
    finally:
        httpd.shutdown()
        httpd.server_close()


def _request(base_url, method, path, body=None):
    data = json.dumps(body).encode("utf-8") if body is not None else None
    request = urllib.request.Request(base_url + path, data=data, method=method)
    if data is not None:
        request.add_header("Content-Type", "application/json")
    try:
        with urllib.request.urlopen(request) as response:
            return response.status, json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        return exc.code, json.loads(exc.read().decode("utf-8"))


def _forecast_body():
    return {
        "starting_portfolio": 500_000,
        "annual_spending": 50_000,
        "annual_retirement_income": 0,
        "nominal_return": 0.05,
        "inflation_rate": 0.02,
        "effective_tax_rate": 0,
        "retirement_start_date": "2030-02-28",
    }


def test_retirement_defaults_endpoint_combines_live_sources(retirement_server):
    base_url, _ = retirement_server
    status, data = _request(base_url, "GET", "/api/retirement/defaults")

    assert status == 200
    assert data["defaults"]["starting_portfolio"] == 500_000
    assert data["defaults"]["annual_spending"] == pytest.approx(3650)
    assert data["defaults"]["nominal_return"] == pytest.approx(0.1383)
    assert data["defaults"]["annual_volatility"] == pytest.approx(0.12)
    assert data["sources"]["budget_as_of"] == "2026-02-28"
    assert data["sources"]["portfolio_as_of"] == "2026-02-28"
    assert data["tax_exempt_portfolio"]["tfsa_value"] == 180_000
    assert data["tax_exempt_portfolio"]["tfsa_share_of_portfolio"] == pytest.approx(0.36)
    assert data["tax_exempt_portfolio"]["non_registered_cost_base"] == 75_000
    assert data["tax_exempt_portfolio"]["non_registered_share_of_portfolio"] == pytest.approx(0.15)
    assert data["tax_exempt_portfolio"]["total_value"] == 255_000
    assert data["tax_exempt_portfolio"]["share_of_portfolio"] == pytest.approx(0.51)


def test_retirement_defaults_endpoint_allows_missing_portfolio(retirement_server):
    base_url, payload_path = retirement_server
    payload_path.unlink()

    status, data = _request(base_url, "GET", "/api/retirement/defaults")

    assert status == 200
    assert data["defaults"]["starting_portfolio"] is None
    assert data["sources"]["portfolio_as_of"] is None


def test_budget_transactions_endpoint_filters_and_returns_row_impact(retirement_server):
    base_url, _ = retirement_server

    status, data = _request(
        base_url,
        "GET",
        "/api/budget/transactions?start=2026-02-01&end=2026-02-28&category=Food&limit=100&offset=0",
    )

    assert status == 200
    assert data == {
        "items": [{
            "row_number": 3,
            "date": "2026-02-28",
            "debit": "Food",
            "credit": "Cash",
            "amount": 590.0,
            "note": "Groceries",
            "impact_cad": 590.0,
        }],
        "total": 1,
        "limit": 100,
        "offset": 0,
    }


def test_retirement_forecast_endpoint_uses_portfolio_as_of_date(retirement_server):
    base_url, _ = retirement_server
    status, data = _request(
        base_url, "POST", "/api/retirement/forecast", _forecast_body()
    )

    assert status == 200
    assert data["start_date"] == "2026-02-28"
    assert data["retirement_start_date"] == "2030-02-28"
    assert data["delay_years"] > 3.9
    assert [row["key"] for row in data["scenarios"]] == [
        "conservative", "base", "optimistic"
    ]
    assert data["probabilistic"]["simulations"] == 10_000


def test_retirement_forecast_endpoint_accepts_probabilistic_schedules(retirement_server):
    base_url, _ = retirement_server
    body = {
        **_forecast_body(),
        "date_of_birth": "1980-02-28",
        "plan_through_age": 55,
        "annual_contribution": 12_000,
        "annual_volatility": 0.11,
        "income_streams": [{
            "name": "Pension", "annual_amount": 10_000,
            "start_date": "2030-02-28", "end_date": None,
        }],
        "spending_phases": [{
            "name": "Core", "annual_amount": 50_000,
            "start_date": "2030-02-28", "end_date": None,
        }],
        "one_time_events": [{
            "name": "Downsize", "date": "2031-02-28", "amount": 75_000,
        }],
    }

    status, data = _request(base_url, "POST", "/api/retirement/forecast", body)

    assert status == 200
    assert data["plan_end_date"] == "2035-02-28"
    assert data["plan_through_age"] == 55
    assert data["probabilistic"]["seed"] == 20260831


def test_retirement_forecast_endpoint_rejects_invalid_input(retirement_server):
    base_url, _ = retirement_server
    body = _forecast_body()
    body["effective_tax_rate"] = 1

    status, data = _request(base_url, "POST", "/api/retirement/forecast", body)

    assert status == 400
    assert "effective_tax_rate" in data["error"]
