"""Tests for the surgical xlsx writer.

Uses a hand-built minimal workbook that mimics the real Personal Budget.xlsx
structure (Journal sheet, Ledger sheet, shared strings, a calcChain part, and
a stand-in "fragile" chart part) so the test suite never depends on the
user's real, gitignored workbook. A skipped-by-default test at the bottom
runs the same checks against the real file when it's present locally.
"""
from __future__ import annotations

import hashlib
import zipfile
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from journal_entry.xlsx_append import (
    JournalLine,
    UnknownAccountError,
    VerificationError,
    WorkbookLockedError,
    append_journal_rows,
    update_journal_row,
)

_STRINGS = [
    "Journal", "Date", "Debit", "Credit", "Amount", "Note",  # 0-5
    "Cash", "Salary", "Pay", "Food", "Groceries",  # 6-10
    "Ledger", "Account", "Beginning Balance",  # 11-13
]

_SHEET1_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheetPr filterMode="1"/><dimension ref="A1:E5"/>
<sheetData>
<row r="1" spans="1:5"><c r="A1" t="s"><v>0</v></c></row>
<row r="2" spans="1:5"><c r="A2" t="s"><v>1</v></c><c r="B2" t="s"><v>2</v></c><c r="C2" t="s"><v>3</v></c><c r="D2" t="s"><v>4</v></c><c r="E2" t="s"><v>5</v></c></row>
<row r="3" spans="1:5"><c r="A3" s="4"><v>46234</v></c><c r="B3" t="s"><v>6</v></c><c r="C3" t="s"><v>7</v></c><c r="D3" s="137"><f>10+5</f><v>15</v></c><c r="E3" t="s"><v>8</v></c></row>
<row r="4" spans="1:5"><c r="A4" s="4"><v>46235</v></c><c r="B4" t="s"><v>9</v></c><c r="C4" t="s"><v>6</v></c><c r="D4" s="137"><v>20</v></c><c r="E4" t="s"><v>10</v></c></row>
<row r="5" spans="1:5"><c r="A5" s="4"/></row>
</sheetData>
<autoFilter ref="A2:E3"/>
<extLst><ext uri="{CCE6A557-97BC-4b89-ADB6-D9C93CAAB3DF}" xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main"><x14:dataValidations count="1" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:dataValidation type="list" allowBlank="1"><x14:formula1><xm:f>Ledger!$A:$A</xm:f></x14:formula1><xm:sqref>B1:C1048576</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>
</worksheet>"""

_SHEET2_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<dimension ref="A1:B5"/>
<sheetData>
<row r="1" spans="1:2"><c r="A1" t="s"><v>11</v></c></row>
<row r="2" spans="1:2"><c r="A2" t="s"><v>12</v></c><c r="B2" t="s"><v>13</v></c></row>
<row r="3" spans="1:2"><c r="A3" t="s"><v>6</v></c><c r="B3"><v>0</v></c></row>
<row r="4" spans="1:2"><c r="A4" t="s"><v>7</v></c><c r="B4"><v>0</v></c></row>
<row r="5" spans="1:2"><c r="A5" t="s"><v>9</v></c><c r="B5"><v>0</v></c></row>
</sheetData>
</worksheet>"""

_SHARED_STRINGS_XML = (
    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
    f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="18" uniqueCount="{len(_STRINGS)}">'
    + "".join(f"<si><t>{s}</t></si>" for s in _STRINGS)
    + "</sst>"
)

_WORKBOOK_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
<sheets>
<sheet name="Journal" sheetId="1" r:id="rId1"/>
<sheet name="Ledger" sheetId="2" r:id="rId2"/>
</sheets>
<calcPr calcId="191029"/>
</workbook>"""

_WORKBOOK_RELS_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
<Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet2.xml"/>
<Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
<Relationship Id="rId4" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/>
<Relationship Id="rId5" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/chart" Target="charts/chart1.xml"/>
<Relationship Id="rId6" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>"""

_CONTENT_TYPES_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/worksheets/sheet2.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/>
<Override PartName="/xl/charts/chart1.xml" ContentType="application/vnd.openxmlformats-officedocument.drawingml.chart+xml"/>
<Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
</Types>"""

_ROOT_RELS_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>"""

_CALC_CHAIN_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<calcChain xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><c r="D3" i="1"/></calcChain>"""

_FRAGILE_CHART_XML = "<chart>FRAGILE-PART-DO-NOT-TOUCH</chart>"

# Style index 4 (dates) needs numFmtId 14 (built-in short date) or openpyxl reads the
# serial as a plain number, not a datetime — like the real workbook's style 4/137.
_STYLE_XFS = "".join(
    f'<xf numFmtId="{14 if i == 4 else 0}" fontId="0" fillId="0" borderId="0" xfId="0"/>'
    for i in range(138)
)
_STYLES_XML = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
<fonts count="1"><font><sz val="11"/><name val="Calibri"/></font></fonts>
<fills count="1"><fill><patternFill patternType="none"/></fill></fills>
<borders count="1"><border><left/><right/><top/><bottom/><diagonal/></border></borders>
<cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
<cellXfs count="138">{_STYLE_XFS}</cellXfs>
</styleSheet>"""


def _build_fixture(path: Path) -> None:
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", _CONTENT_TYPES_XML)
        zf.writestr("_rels/.rels", _ROOT_RELS_XML)
        zf.writestr("xl/workbook.xml", _WORKBOOK_XML)
        zf.writestr("xl/_rels/workbook.xml.rels", _WORKBOOK_RELS_XML)
        zf.writestr("xl/worksheets/sheet1.xml", _SHEET1_XML)
        zf.writestr("xl/worksheets/sheet2.xml", _SHEET2_XML)
        zf.writestr("xl/sharedStrings.xml", _SHARED_STRINGS_XML)
        zf.writestr("xl/styles.xml", _STYLES_XML)
        zf.writestr("xl/calcChain.xml", _CALC_CHAIN_XML)
        zf.writestr("xl/charts/chart1.xml", _FRAGILE_CHART_XML)


@pytest.fixture
def workbook(tmp_path):
    path = tmp_path / "Personal Budget.xlsx"
    _build_fixture(path)
    return path


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_appends_rows_readable_via_openpyxl(workbook):
    lines = [
        JournalLine(date(2026, 8, 30), "Food", "Cash", [12.5, 8.0], "Groceries"),
        JournalLine(date(2026, 8, 30), "Cash", "Salary", [1000.0], "Pay"),
    ]
    result = append_journal_rows(workbook, lines)
    assert result.first_row == 6
    assert result.last_row == 7

    wb = load_workbook(workbook, read_only=True, data_only=True)
    rows = list(wb["Journal"].iter_rows(min_row=6, max_row=7, max_col=5, values_only=True))
    wb.close()
    assert rows[0][1:] == ("Food", "Cash", 20.5, "Groceries")
    assert rows[1][1:] == ("Cash", "Salary", 1000.0, "Pay")


def test_multiple_new_shared_strings_in_one_batch_keep_sequential_indexes(workbook):
    lines = [
        JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "First new note"),
        JournalLine(date(2026, 8, 30), "Cash", "Salary", [7.0], "Second new note"),
    ]

    result = append_journal_rows(workbook, lines)

    wb = load_workbook(workbook, read_only=True, data_only=True)
    rows = list(wb["Journal"].iter_rows(
        min_row=result.first_row, max_row=result.last_row, max_col=5, values_only=True
    ))
    wb.close()
    assert rows[0][4] == "First new note"
    assert rows[1][4] == "Second new note"


def test_single_component_writes_plain_value_no_formula(workbook):
    append_journal_rows(workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [42.0], "Snack")])
    with zipfile.ZipFile(workbook) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<c r="D6" s="137"><v>42</v></c>' in sheet


def test_multi_component_writes_sum_formula(workbook):
    append_journal_rows(
        workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [12.5, 8.0], "Groceries")]
    )
    with zipfile.ZipFile(workbook) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert "<f>12.5+8</f><v>20.5</v>" in sheet


def test_reuses_existing_shared_strings_only_adds_new_note(workbook):
    with zipfile.ZipFile(workbook) as zf:
        before_unique = int(
            zf.read("xl/sharedStrings.xml").decode("utf-8").split('uniqueCount="')[1].split('"')[0]
        )
    append_journal_rows(
        workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "Brand new note text")]
    )
    with zipfile.ZipFile(workbook) as zf:
        ss = zf.read("xl/sharedStrings.xml").decode("utf-8")
    after_unique = int(ss.split('uniqueCount="')[1].split('"')[0])
    assert after_unique == before_unique + 1
    assert "Brand new note text" in ss


def test_multiple_new_notes_in_one_call_get_distinct_valid_indices(workbook):
    append_journal_rows(workbook, [
        JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "first new note"),
        JournalLine(date(2026, 8, 31), "Cash", "Salary", [10.0], "second new note"),
    ])
    wb = load_workbook(workbook, read_only=True, data_only=True)
    rows = list(wb["Journal"].iter_rows(min_row=6, max_row=7, max_col=5, values_only=True))
    wb.close()
    assert rows[0][4] == "first new note"
    assert rows[1][4] == "second new note"


def test_dimension_updated(workbook):
    append_journal_rows(workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "x")])
    with zipfile.ZipFile(workbook) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<dimension ref="A1:E6"/>' in sheet


def test_fragile_parts_untouched_byte_for_byte(workbook):
    with zipfile.ZipFile(workbook) as zf:
        before_chart = zf.read("xl/charts/chart1.xml")
        before_sheet2 = zf.read("xl/worksheets/sheet2.xml")
    append_journal_rows(workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "x")])
    with zipfile.ZipFile(workbook) as zf:
        assert zf.read("xl/charts/chart1.xml") == before_chart
        assert zf.read("xl/worksheets/sheet2.xml") == before_sheet2
        # x14 dataValidation / autoFilter untouched — the user's Excel view is not disturbed
        sheet1 = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert '<autoFilter ref="A2:E3"/>' in sheet1
        assert "Ledger!$A:$A" in sheet1


def test_calc_chain_removed_cleanly(workbook):
    append_journal_rows(workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "x")])
    with zipfile.ZipFile(workbook) as zf:
        names = zf.namelist()
        assert "xl/calcChain.xml" not in names
        assert "calcChain" not in zf.read("[Content_Types].xml").decode("utf-8")
        assert "calcChain" not in zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        assert 'fullCalcOnLoad="1"' in zf.read("xl/workbook.xml").decode("utf-8")


def test_backup_created_before_write(workbook):
    result = append_journal_rows(workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "x")])
    assert result.backup_path.is_file()
    backup_wb = load_workbook(result.backup_path, read_only=True, data_only=True)
    # backup reflects the pre-write state: only the original 2 data rows
    assert backup_wb["Journal"].max_row == 5
    backup_wb.close()


def test_unknown_account_rejected_without_touching_workbook(workbook):
    before = _hash(workbook)
    with pytest.raises(UnknownAccountError):
        append_journal_rows(
            workbook, [JournalLine(date(2026, 8, 30), "Nonexistent Category", "Cash", [5.0], "x")]
        )
    assert _hash(workbook) == before


def test_locked_workbook_refused(workbook):
    lock_file = workbook.with_name(f"~${workbook.name}")
    lock_file.write_text("lock")
    with pytest.raises(WorkbookLockedError):
        append_journal_rows(workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "x")])


def test_empty_lines_rejected(workbook):
    with pytest.raises(ValueError):
        append_journal_rows(workbook, [])


def test_second_batch_appends_after_first(workbook):
    append_journal_rows(workbook, [JournalLine(date(2026, 8, 30), "Food", "Cash", [5.0], "first")])
    result = append_journal_rows(workbook, [JournalLine(date(2026, 8, 31), "Food", "Cash", [7.0], "second")])
    assert result.first_row == 7
    wb = load_workbook(workbook, read_only=True, data_only=True)
    assert wb["Journal"].max_row == 7
    wb.close()


def test_update_journal_row_overwrites_in_place(workbook):
    backup_path = update_journal_row(
        workbook, 4, JournalLine(date(2026, 9, 1), "Food", "Cash", [99.0], "Corrected")
    )
    assert backup_path.is_file()

    wb = load_workbook(workbook, read_only=True, data_only=True)
    row = list(wb["Journal"].iter_rows(min_row=4, max_row=4, max_col=5, values_only=True))[0]
    wb.close()
    assert row[1:] == ("Food", "Cash", 99.0, "Corrected")


def test_update_journal_row_does_not_change_dimension_or_other_rows(workbook):
    update_journal_row(workbook, 4, JournalLine(date(2026, 9, 1), "Food", "Cash", [99.0], "Corrected"))
    with zipfile.ZipFile(workbook) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<dimension ref="A1:E5"/>' in sheet
    # Row 3 (untouched) still has its original formula.
    assert "<f>10+5</f><v>15</v>" in sheet


def test_update_journal_row_flattens_formula_to_plain_value(workbook):
    update_journal_row(workbook, 3, JournalLine(date(2026, 9, 1), "Food", "Cash", [7.0], "Fixed"))
    with zipfile.ZipFile(workbook) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    assert '<c r="D3" s="137"><v>7</v></c>' in sheet
    assert "<f>10+5</f>" not in sheet


def test_update_journal_row_unknown_account_rejected_without_touching_workbook(workbook):
    before = _hash(workbook)
    with pytest.raises(UnknownAccountError):
        update_journal_row(
            workbook, 4, JournalLine(date(2026, 9, 1), "Nonexistent Category", "Cash", [7.0], "x")
        )
    assert _hash(workbook) == before


def test_update_journal_row_locked_workbook_refused(workbook):
    lock_file = workbook.with_name(f"~${workbook.name}")
    lock_file.write_text("lock")
    with pytest.raises(WorkbookLockedError):
        update_journal_row(workbook, 4, JournalLine(date(2026, 9, 1), "Food", "Cash", [7.0], "x"))


def test_update_journal_row_missing_row_rejected_without_touching_workbook(workbook):
    before = _hash(workbook)
    with pytest.raises(ValueError):
        update_journal_row(workbook, 999, JournalLine(date(2026, 9, 1), "Food", "Cash", [7.0], "x"))
    assert _hash(workbook) == before


def test_update_journal_row_self_closing_target_does_not_delete_next_row(workbook):
    with zipfile.ZipFile(workbook) as zf:
        sheet = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
    # Rewrite row 3 (normally <row r="3" spans="1:5"><c.../></row>) as
    # self-closing and empty -- e.g. a hand-cleared cell in Excel -- while
    # leaving row 4 (Debit=Food, Credit=Cash, Amount=20, Note=Groceries)
    # as a normal row right after it.
    original_row3 = (
        '<row r="3" spans="1:5"><c r="A3" s="4"><v>46234</v></c>'
        '<c r="B3" t="s"><v>6</v></c><c r="C3" t="s"><v>7</v></c>'
        '<c r="D3" s="137"><f>10+5</f><v>15</v></c><c r="E3" t="s"><v>8</v></c></row>'
    )
    assert original_row3 in sheet
    patched_sheet = sheet.replace(original_row3, '<row r="3" spans="1:5"/>', 1)
    with zipfile.ZipFile(workbook, "r") as src:
        items = {item.filename: src.read(item.filename) for item in src.infolist()}
    items["xl/worksheets/sheet1.xml"] = patched_sheet.encode("utf-8")
    with zipfile.ZipFile(workbook, "w", zipfile.ZIP_DEFLATED) as dst:
        for name, data in items.items():
            dst.writestr(name, data)

    update_journal_row(workbook, 3, JournalLine(date(2026, 9, 1), "Food", "Cash", [7.0], "Fixed"))

    wb = load_workbook(workbook, read_only=True, data_only=True)
    row3 = list(wb["Journal"].iter_rows(min_row=3, max_row=3, max_col=5, values_only=True))[0]
    row4 = list(wb["Journal"].iter_rows(min_row=4, max_row=4, max_col=5, values_only=True))[0]
    wb.close()
    assert row3[1:] == ("Food", "Cash", 7.0, "Fixed")
    assert row4[1:] == ("Food", "Cash", 20, "Groceries")  # row 4 must survive untouched


def test_update_journal_row_backup_reflects_pre_edit_value(workbook):
    backup_path = update_journal_row(
        workbook, 4, JournalLine(date(2026, 9, 1), "Food", "Cash", [7.0], "Fixed")
    )
    backup_wb = load_workbook(backup_path, read_only=True, data_only=True)
    row = list(backup_wb["Journal"].iter_rows(min_row=4, max_row=4, max_col=5, values_only=True))[0]
    backup_wb.close()
    assert row[3] == 20  # original row 4 amount from the fixture, before this edit


def test_update_journal_row_fragile_parts_untouched(workbook):
    with zipfile.ZipFile(workbook) as zf:
        before_chart = zf.read("xl/charts/chart1.xml")
        before_sheet2 = zf.read("xl/worksheets/sheet2.xml")
    update_journal_row(workbook, 4, JournalLine(date(2026, 9, 1), "Food", "Cash", [7.0], "x"))
    with zipfile.ZipFile(workbook) as zf:
        assert zf.read("xl/charts/chart1.xml") == before_chart
        assert zf.read("xl/worksheets/sheet2.xml") == before_sheet2
        sheet1 = zf.read("xl/worksheets/sheet1.xml").decode("utf-8")
        assert '<autoFilter ref="A2:E3"/>' in sheet1
        assert "Ledger!$A:$A" in sheet1


def test_update_journal_row_calc_chain_removed_cleanly(workbook):
    update_journal_row(workbook, 4, JournalLine(date(2026, 9, 1), "Food", "Cash", [7.0], "x"))
    with zipfile.ZipFile(workbook) as zf:
        names = zf.namelist()
        assert "xl/calcChain.xml" not in names
        assert "calcChain" not in zf.read("[Content_Types].xml").decode("utf-8")
        assert "calcChain" not in zf.read("xl/_rels/workbook.xml.rels").decode("utf-8")
        assert 'fullCalcOnLoad="1"' in zf.read("xl/workbook.xml").decode("utf-8")


_REAL_WORKBOOK = Path(__file__).resolve().parents[1] / "data" / "Personal Budget.xlsx"


@pytest.mark.skipif(not _REAL_WORKBOOK.is_file(), reason="real workbook not present on this machine")
def test_round_trip_against_real_workbook_preserves_fragile_parts(tmp_path):
    import shutil

    copy_path = tmp_path / "Personal Budget.xlsx"
    shutil.copy2(_REAL_WORKBOOK, copy_path)

    with zipfile.ZipFile(_REAL_WORKBOOK) as before:
        before_names = set(before.namelist())
        before_bytes = {n: before.read(n) for n in before_names}

    result = append_journal_rows(
        copy_path, [JournalLine(date(2026, 8, 30), "Food", "Cash", [1.23], "journal-entry test row")]
    )
    assert result.backup_path.is_file()

    with zipfile.ZipFile(copy_path) as after:
        after_names = set(after.namelist())
        changed = {
            n for n in before_names & after_names if before_bytes[n] != after.read(n)
        }

    assert before_names - after_names == {"xl/calcChain.xml"}
    assert changed == {
        "xl/worksheets/sheet1.xml",
        "xl/sharedStrings.xml",
        "xl/workbook.xml",
        "[Content_Types].xml",
        "xl/_rels/workbook.xml.rels",
    }
