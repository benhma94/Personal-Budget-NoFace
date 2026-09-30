"""Month-end balance reconciliations for the Journal workflow.

The two supported adjustments use the same accounting rule: compare an
actual ending balance with the balance implied by the workbook and the current
Post queue, then stage the difference as a signed transaction.  Positive
differences debit the asset; negative differences credit it.
"""
from __future__ import annotations

import calendar
import re
import uuid
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from budget_dashboard.profile import PROFILE

from .aggregate import build_posting_lines
from .store import Store

_CENT = Decimal("0.01")
_MONTH_RE = re.compile(r"^(\d{4})-(\d{2})$")
_ACTIVE_STATUSES = ("new", "categorized")

_CLOSE_TYPES = PROFILE["month_close"]


def closing_date(month: str) -> date:
    """Return the calendar month-end represented by ``YYYY-MM``."""
    match = _MONTH_RE.fullmatch(str(month or ""))
    if not match:
        raise ValueError("month must use YYYY-MM format")
    year, month_number = map(int, match.groups())
    if not 1 <= month_number <= 12:
        raise ValueError("month must use YYYY-MM format")
    return date(year, month_number, calendar.monthrange(year, month_number)[1])


def _money(value: Any, field: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError(f"{field} is required")
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError):
        raise ValueError(f"{field} must be a valid amount") from None
    if not result.is_finite():
        raise ValueError(f"{field} must be a finite amount")
    result = result.quantize(_CENT, rounding=ROUND_HALF_UP)
    if result < 0:
        raise ValueError(f"{field} cannot be negative")
    return result


def _as_money(value: Any) -> Decimal:
    """Convert trusted workbook/queue numbers to currency precision."""
    try:
        result = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return Decimal("0.00")
    if not result.is_finite():
        return Decimal("0.00")
    return result.quantize(_CENT, rounding=ROUND_HALF_UP)


def _number(value: Decimal | None) -> float | None:
    return float(value) if value is not None else None


def _workbook_balances(
    workbook_path: str | Path, cutoff: date
) -> tuple[dict[str, str], dict[str, Decimal]]:
    """Resolve closing accounts and calculate their posted balances at cutoff."""
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if "Journal" not in wb.sheetnames or "Ledger" not in wb.sheetnames:
            raise ValueError("workbook must contain Journal and Ledger sheets")

        ledger = wb["Ledger"]
        ledger_rows = [
            (str(account).strip(), _as_money(opening))
            for account, opening in ledger.iter_rows(min_row=3, max_col=2, values_only=True)
            if account not in (None, "")
        ]
        by_casefold = {name.casefold(): name for name, _ in ledger_rows}
        required = {
            value.casefold()
            for config in _CLOSE_TYPES.values()
            for value in (config["account"], config["category"])
        }
        missing = sorted(value for value in required if value not in by_casefold)
        if missing:
            raise ValueError(
                "Ledger is missing month-close account(s): " + ", ".join(missing)
            )

        resolved = {value: by_casefold[value] for value in required}
        balance_names = {
            resolved[config["account"].casefold()] for config in _CLOSE_TYPES.values()
        }
        balances = {
            name: opening for name, opening in ledger_rows if name in balance_names
        }

        journal = wb["Journal"]
        for raw_date, debit, credit, amount, *_ in journal.iter_rows(
            min_row=3, max_col=5, values_only=True
        ):
            if not isinstance(raw_date, (date, datetime)):
                continue
            txn_date = raw_date.date() if isinstance(raw_date, datetime) else raw_date
            if txn_date > cutoff:
                continue
            value = _as_money(amount)
            debit_name = resolved.get(str(debit or "").strip().casefold())
            credit_name = resolved.get(str(credit or "").strip().casefold())
            if debit_name in balances:
                balances[debit_name] += value
            if credit_name in balances:
                balances[credit_name] -= value

        return resolved, {
            name: value.quantize(_CENT, rounding=ROUND_HALF_UP)
            for name, value in balances.items()
        }
    finally:
        wb.close()


def _txn_date(row: dict[str, Any]) -> date:
    try:
        return date.fromisoformat(str(row["txn_date"]))
    except (KeyError, TypeError, ValueError):
        raise ValueError("queued transaction has an invalid transaction date") from None


def _portfolio_info(
    portfolio_snapshot: dict[str, Any] | None, cutoff: date
) -> tuple[dict[str, Any], Decimal | None, list[str]]:
    snapshot = portfolio_snapshot or {}
    raw_balance = snapshot.get("total_value_cad")
    balance: Decimal | None = None
    if raw_balance is not None:
        try:
            candidate = Decimal(str(raw_balance))
            if candidate.is_finite() and candidate >= 0:
                balance = candidate.quantize(_CENT, rounding=ROUND_HALF_UP)
        except (InvalidOperation, ValueError):
            pass

    as_of = snapshot.get("as_of_date")
    generated_at = snapshot.get("generated_at")
    warnings: list[str] = []
    if balance is None:
        warnings.append(
            "No cached portfolio total is available; enter the ending portfolio balance manually."
        )
    if balance is not None and as_of != cutoff.isoformat():
        label = as_of or "an unknown date"
        warnings.append(
            f"The cached portfolio is as of {label}, not {cutoff.isoformat()}; verify or override it."
        )
    return {
        "total_value_cad": _number(balance),
        "as_of_date": as_of,
        "generated_at": generated_at,
    }, balance, warnings


def build_preview(
    workbook_path: str | Path,
    store: Store,
    confirmed_transfers: list[dict[str, Any]],
    month: str,
    portfolio_snapshot: dict[str, Any] | None,
    selection: dict[str, Any] | None = None,
) -> dict[str, Any]:
    """Build the month-close calculation without changing staged state."""
    cutoff = closing_date(month)
    resolved, workbook_balances = _workbook_balances(workbook_path, cutoff)
    cache_info, cached_portfolio, warnings = _portfolio_info(portfolio_snapshot, cutoff)

    all_categorized = store.list_transactions(status="categorized")
    normal_categorized = [
        row for row in all_categorized if row.get("origin") != "month_close"
    ]
    eligible_categorized = [
        row for row in normal_categorized if _txn_date(row) <= cutoff
    ]
    future_fingerprints = {
        row["fingerprint"] for row in normal_categorized if _txn_date(row) > cutoff
    }

    eligible_transfers: list[dict[str, Any]] = []
    for transfer in confirmed_transfers:
        rows = (transfer["outgoing"], transfer["incoming"])
        if any(_txn_date(row) > cutoff for row in rows):
            future_fingerprints.update(row["fingerprint"] for row in rows)
        else:
            eligible_transfers.append(transfer)

    lines, _ = build_posting_lines(eligible_categorized, eligible_transfers, cutoff)
    pending_effects = {name: Decimal("0.00") for name in workbook_balances}
    target_by_casefold = {name.casefold(): name for name in workbook_balances}
    for line in lines:
        amount = _as_money(line.amount)
        debit = target_by_casefold.get(line.debit.casefold())
        credit = target_by_casefold.get(line.credit.casefold())
        if debit:
            pending_effects[debit] += amount
        if credit:
            pending_effects[credit] -= amount

    new_rows = store.list_transactions(status="new")
    other_month_rows = [
        row
        for row in store.list_month_close_transactions(statuses=_ACTIVE_STATUSES)
        if row.get("close_month") != month
    ]
    blockers: list[str] = []
    if new_rows:
        blockers.append(
            f"Resolve or ignore {len(new_rows)} unconfirmed Review transaction(s) before staging."
        )
    if future_fingerprints:
        blockers.append(
            f"Resolve {len(future_fingerprints)} postable transaction(s) dated after {cutoff.isoformat()}."
        )
    if other_month_rows:
        other_months = ", ".join(sorted({row["close_month"] for row in other_month_rows}))
        blockers.append(
            f"Post or remove the staged month-close adjustment(s) for {other_months} first."
        )

    active_rows = store.list_month_close_transactions(
        close_month=month, statuses=_ACTIVE_STATUSES
    )
    active_by_kind = {row["close_kind"]: row for row in active_rows}

    projected: dict[str, Decimal] = {}
    for kind, config in _CLOSE_TYPES.items():
        account = resolved[config["account"].casefold()]
        projected[kind] = (
            workbook_balances[account] + pending_effects[account]
        ).quantize(_CENT, rounding=ROUND_HALF_UP)

    if selection is None:
        selection = {}
        for kind in _CLOSE_TYPES:
            staged = active_by_kind.get(kind)
            if staged:
                target = projected[kind] + _as_money(staged["amount"])
                selection[kind] = {"enabled": True, "ending_balance": target}
            elif kind == "investment" and cached_portfolio is not None:
                selection[kind] = {
                    "enabled": True,
                    "ending_balance": cached_portfolio,
                }
            else:
                selection[kind] = {"enabled": False, "ending_balance": None}

    reconciliations: dict[str, dict[str, Any]] = {}
    for kind, config in _CLOSE_TYPES.items():
        spec = selection.get(kind) or {}
        enabled = spec.get("enabled") is True
        target = _money(spec.get("ending_balance"), f"{config['label']} ending balance") if enabled else None
        adjustment = (
            (target - projected[kind]).quantize(_CENT, rounding=ROUND_HALF_UP)
            if target is not None
            else None
        )
        account = resolved[config["account"].casefold()]
        category = resolved[config["category"].casefold()]
        debit = credit = None
        amount = None
        if adjustment is not None and adjustment != 0:
            debit, credit = (
                (account, category) if adjustment > 0 else (category, account)
            )
            amount = abs(adjustment)
        reconciliations[kind] = {
            "kind": kind,
            "label": config["label"],
            "enabled": enabled,
            "account": account,
            "category": category,
            "workbook_balance": _number(workbook_balances[account]),
            "pending_effect": _number(pending_effects[account]),
            "projected_balance": _number(projected[kind]),
            "ending_balance": _number(target),
            "adjustment": _number(adjustment),
            "debit": debit,
            "credit": credit,
            "amount": _number(amount),
            "note": config["note"],
            "staged": kind in active_by_kind,
        }

    has_selection = any(row["enabled"] for row in reconciliations.values())
    can_clear = bool(active_rows)
    return {
        "month": month,
        "closing_date": cutoff.isoformat(),
        "portfolio_cache": cache_info,
        "warnings": warnings,
        "blockers": blockers,
        "reconciliations": reconciliations,
        "can_stage": not blockers and (has_selection or can_clear),
    }


def stage_adjustments(
    workbook_path: str | Path,
    store: Store,
    confirmed_transfers: list[dict[str, Any]],
    month: str,
    portfolio_snapshot: dict[str, Any] | None,
    selection: dict[str, Any],
) -> dict[str, Any]:
    """Replace this month's active close rows with the selected adjustments."""
    preview = build_preview(
        workbook_path, store, confirmed_transfers, month, portfolio_snapshot, selection
    )
    if preview["blockers"]:
        raise ValueError(" ".join(preview["blockers"]))

    now = datetime.now().isoformat(timespec="seconds")
    rows_by_kind: dict[str, dict[str, Any] | None] = {}
    for kind, config in _CLOSE_TYPES.items():
        reconciliation = preview["reconciliations"][kind]
        adjustment = reconciliation["adjustment"]
        if not reconciliation["enabled"] or adjustment in (None, 0):
            rows_by_kind[kind] = None
            continue
        rows_by_kind[kind] = {
            "fingerprint": f"month-close:{month}:{kind}:{uuid.uuid4().hex[:12]}",
            "account": reconciliation["account"],
            "txn_date": preview["closing_date"],
            "description": config["description"],
            "amount": adjustment,
            "category": reconciliation["category"],
            "note": config["note"],
            "status": "categorized",
            "transfer_peer": None,
            "batch_id": None,
            "imported_at": now,
            "close_month": month,
            "close_kind": kind,
        }

    store.replace_month_close_transactions(month, rows_by_kind)
    result = build_preview(
        workbook_path, store, confirmed_transfers, month, portfolio_snapshot, selection
    )
    result["staged_count"] = sum(
        1 for row in result["reconciliations"].values() if row["staged"]
    )
    return result
