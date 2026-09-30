"""Read specific Journal-sheet rows without modifying the workbook. Mirrors
the read-only pattern budget_dashboard/workbook.py already uses.
"""
from __future__ import annotations

import warnings
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# OpenPyXL cannot preserve an extended validation object, but this reader
# never saves the workbook, so the warning does not apply here.
warnings.filterwarnings(
    "ignore",
    message="Data Validation extension is not supported.*",
    module="openpyxl.worksheet._reader",
)


def read_journal_rows(workbook_path: str | Path, first_row: int, last_row: int) -> list[dict[str, Any]]:
    """Return Journal!A{first_row}:E{last_row} as of the workbook's last
    save, including whether each Amount cell holds a formula (a row
    aggregated from multiple imported transactions) rather than a plain
    value -- that distinction isn't visible through a data-only read, so
    the sheet is opened twice: once for cached values, once for formulas.
    """
    workbook_path = Path(workbook_path)
    values_wb = None
    formulas_wb = None
    try:
        values_wb = load_workbook(workbook_path, read_only=True, data_only=True)
        formulas_wb = load_workbook(workbook_path, read_only=True, data_only=False)
        value_rows = values_wb["Journal"].iter_rows(min_row=first_row, max_row=last_row, max_col=5)
        formula_rows = formulas_wb["Journal"].iter_rows(min_row=first_row, max_row=last_row, max_col=5)
        rows = []
        for row_number, (value_row, formula_row) in enumerate(zip(value_rows, formula_rows), start=first_row):
            posting_date, debit, credit, amount, note = (cell.value for cell in value_row)
            rows.append({
                "row_number": row_number,
                "posting_date": posting_date.strftime("%Y-%m-%d") if isinstance(posting_date, (date, datetime))
                    else (str(posting_date) if posting_date is not None else None),
                "debit": debit,
                "credit": credit,
                "amount": float(amount) if isinstance(amount, (int, float)) else None,
                "note": note,
                "is_formula": formula_row[3].data_type == "f",
            })
        return rows
    finally:
        if values_wb is not None:
            values_wb.close()
        if formulas_wb is not None:
            formulas_wb.close()
