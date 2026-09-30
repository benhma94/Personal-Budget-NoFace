"""Read price overrides and write Report_* sheets to the user's workbook.

Transactions are loaded from a Wealthsimple CSV (see io_csv.py), not from
the workbook. User-maintained sheets are never modified. Only `Report_*`
sheets are overwritten on each run.
"""
from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Iterable

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill


_REPORT_PREFIX = "Report_"


def read_instrument_map(workbook_path) -> dict[str, dict[str, str]]:
    """Read optional canonical ticker -> provider identifier mappings."""
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    if "InstrumentMap" not in wb.sheetnames:
        return {}
    rows = list(wb["InstrumentMap"].iter_rows(values_only=True))
    if not rows:
        return {}
    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    required = {"ticker", "lseg_ric", "yahoo_ticker"}
    if not required.issubset(header):
        raise ValueError("InstrumentMap requires ticker, lseg_ric, and yahoo_ticker columns")
    idx = {name: header.index(name) for name in header if name}
    result: dict[str, dict[str, str]] = {}
    for row in rows[1:]:
        value = row[idx["ticker"]] if len(row) > idx["ticker"] else None
        if value in (None, ""):
            continue
        ticker = str(value).strip()
        if ticker in result:
            raise ValueError(f"InstrumentMap contains duplicate ticker {ticker!r}")
        result[ticker] = {
            "lseg_ric": str(row[idx["lseg_ric"]]).strip() if len(row) > idx["lseg_ric"] and row[idx["lseg_ric"]] else ticker,
            "yahoo_ticker": str(row[idx["yahoo_ticker"]]).strip() if len(row) > idx["yahoo_ticker"] and row[idx["yahoo_ticker"]] else ticker,
            "note": str(row[idx["note"]]).strip() if "note" in idx and len(row) > idx["note"] and row[idx["note"]] else "",
        }
    return result


def read_price_overrides(workbook_path) -> dict[str, dict]:
    """Read the optional PriceOverrides sheet.

    Sheet layout:
        ticker | currency | price | note

    Returns {ticker: {currency, price, note}}. Empty dict if the sheet is
    missing — overrides are entirely optional.
    """
    wb = load_workbook(workbook_path, data_only=True, read_only=True)
    if "PriceOverrides" not in wb.sheetnames:
        return {}

    ws = wb["PriceOverrides"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return {}

    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    required = {"ticker", "currency", "price"}
    if not required.issubset(set(header)):
        return {}

    idx = {h: header.index(h) for h in header if h}
    out = {}
    for row in rows[1:]:
        if not row or row[idx["ticker"]] in (None, ""):
            continue
        ticker = str(row[idx["ticker"]]).strip()
        out[ticker] = {
            "currency": (str(row[idx["currency"]]) if row[idx["currency"]] else "CAD").strip().upper(),
            "price": float(row[idx["price"]]),
            "note": str(row[idx["note"]]) if "note" in idx and row[idx["note"]] else "",
        }
    return out


# ----- writing reports -----

@dataclass
class ReportTables:
    positions: pd.DataFrame
    geography: pd.DataFrame
    sector: pd.DataFrame
    asset_class: pd.DataFrame
    performance: pd.DataFrame
    risk: pd.DataFrame
    meta: dict


def write_reports(workbook_path, tables: ReportTables) -> None:
    wb = load_workbook(workbook_path)

    # Drop existing report sheets (overwrite, not append)
    for name in list(wb.sheetnames):
        if name.startswith(_REPORT_PREFIX):
            del wb[name]

    _write_positions(wb, tables.positions)
    _write_exposure(wb, tables.geography, tables.sector, tables.asset_class)
    _write_performance(wb, tables.performance)
    _write_risk(wb, tables.risk)
    _write_meta(wb, tables.meta)

    wb.save(workbook_path)


def _header_font():
    return Font(bold=True, color="FFFFFF")


def _header_fill():
    return PatternFill("solid", fgColor="305496")


def _write_dataframe(ws, df: pd.DataFrame, start_row: int = 1, title: str | None = None) -> int:
    row = start_row
    if title:
        cell = ws.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, size=12)
        row += 1

    for col_idx, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=row, column=col_idx, value=str(name))
        cell.font = _header_font()
        cell.fill = _header_fill()
    row += 1

    for _, record in df.iterrows():
        for col_idx, name in enumerate(df.columns, start=1):
            value = record[name]
            if hasattr(value, "item"):
                value = value.item()
            ws.cell(row=row, column=col_idx, value=value)
        row += 1

    return row + 1  # leave one blank row after


def _write_positions(wb, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(f"{_REPORT_PREFIX}Positions")
    _write_dataframe(ws, df)
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = 18


def _write_exposure(wb, geo: pd.DataFrame, sec: pd.DataFrame, ac: pd.DataFrame) -> None:
    ws = wb.create_sheet(f"{_REPORT_PREFIX}Exposure")
    next_row = _write_dataframe(ws, geo, start_row=1, title="Geographic Exposure")
    next_row = _write_dataframe(ws, sec, start_row=next_row, title="Sector Exposure")
    _write_dataframe(ws, ac, start_row=next_row, title="Asset Class Exposure")
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 14


def _write_performance(wb, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(f"{_REPORT_PREFIX}Performance")
    _write_dataframe(ws, df)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    for col_letter in ["C", "D", "E"]:
        ws.column_dimensions[col_letter].width = 16
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=5):
        for cell in row:
            cell.number_format = "0.00%"


def _write_risk(wb, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(f"{_REPORT_PREFIX}Risk")
    _write_dataframe(ws, df)
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 18


def _write_meta(wb, meta: dict) -> None:
    ws = wb.create_sheet(f"{_REPORT_PREFIX}Meta")
    df = pd.DataFrame([(k, v) for k, v in meta.items()], columns=["key", "value"])
    _write_dataframe(ws, df)
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 60
